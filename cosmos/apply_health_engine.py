import pandas as pd
import numpy as np
import joblib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import os

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.abspath(os.path.join(BASE_DIR, ".."))

SRC = os.path.join(ROOT_DIR, "cosmos_twin_dataset.xlsx")
if not os.path.exists(SRC):
    SRC = os.path.join(BASE_DIR, "cosmos_twin_dataset.xlsx")

OUT = os.path.join(BASE_DIR, "cosmos_twin_ml_output_v2.xlsx")
MODEL_PATH = os.path.join(BASE_DIR, "models.joblib")

bundle = joblib.load(MODEL_PATH)
FEATURES = bundle["features"]

tel = pd.read_excel(SRC, sheet_name="Telemetry").sort_values("timestamp").reset_index(drop=True)
tel["payload_on"] = (tel["payload_status"] == "ON").astype(int)
tel["comm_linked"] = (tel["communication_status"] == "LINKED").astype(int)
tel["sunlit_flag"] = tel["sunlit"].astype(int)

BASE_SIGNALS = ["battery_percentage", "battery_voltage", "battery_current",
                "temperature", "solar_power_input", "power_draw"]
for col in BASE_SIGNALS:
    tel[f"{col}_roll5_mean"] = tel.groupby("satellite_id")[col].transform(lambda s: s.rolling(5, min_periods=1).mean())
    tel[f"{col}_roll5_std"] = tel.groupby("satellite_id")[col].transform(lambda s: s.rolling(5, min_periods=1).std().fillna(0))
    tel[f"{col}_roll10_mean"] = tel.groupby("satellite_id")[col].transform(lambda s: s.rolling(10, min_periods=1).mean())
    tel[f"{col}_delta"] = tel.groupby("satellite_id")[col].diff().fillna(0)

X = tel[FEATURES]

tel["battery_health"] = bundle["battery"].predict(X).clip(0, 100)
tel["thermal_health"] = bundle["thermal"].predict(X).clip(0, 100)
tel["power_health"] = bundle["power"].predict(X).clip(0, 100)
tel["payload_health"] = bundle["payload"].predict(X).clip(0, 100)
tel["health_score"] = bundle["health"].predict(X).clip(0, 100)

anom_proba = bundle["anomaly"].predict_proba(X)[:, 1]
tel["anomaly_flag"] = (anom_proba >= 0.5).astype(bool)

chan_probs = {chan: m.predict_proba(X)[:, 1] for chan, m in bundle["chan_models"].items()}
chan_df = pd.DataFrame(chan_probs)
tel["anomaly_parameter"] = np.where(tel["anomaly_flag"], chan_df.idxmax(axis=1), "none")
tel["anomaly_confidence"] = np.where(tel["anomaly_flag"], chan_df.max(axis=1).round(3), 0.0)

def risk(score):
    if score < 40: return "CRITICAL"
    if score < 60: return "HIGH"
    if score < 80: return "MED"
    return "LOW"
tel["risk_level"] = tel["health_score"].round(2).apply(risk)

tel["twin_state_id"] = [f"TWIN-{i:06d}" for i in range(len(tel))]
tel["source_row_index"] = range(len(tel))
tel["sync_status"] = "LIVE_REPLAY"
tel["state_confidence"] = 1.0

for c in ["health_score","battery_health","thermal_health","power_health","payload_health"]:
    tel[c] = tel[c].round(2)

drop_cols = ["payload_on","comm_linked","sunlit_flag"] + \
    [f"{c}_roll5_mean" for c in BASE_SIGNALS] + [f"{c}_roll5_std" for c in BASE_SIGNALS] + \
    [f"{c}_roll10_mean" for c in BASE_SIGNALS] + [f"{c}_delta" for c in BASE_SIGNALS]
tel_out = tel.drop(columns=drop_cols)

wb = Workbook()
ws = wb.active
ws.title = "ML_Twin_Output"
header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
body_font = Font(name="Arial")

for c, col in enumerate(tel_out.columns, start=1):
    cell = ws.cell(row=1, column=c, value=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")
for r, row in enumerate(tel_out.itertuples(index=False), start=2):
    for c, val in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=val).font = body_font
for c, col in enumerate(tel_out.columns, start=1):
    width = max(12, min(22, int(tel_out[col].astype(str).str.len().max()) + 2))
    ws.column_dimensions[get_column_letter(c)].width = width
ws.freeze_panes = "A2"

ws2 = wb.create_sheet("Model_Performance")
rows = [
    ("Model","Type","Metric","Before","After","Notes"),
    ("health_score","XGBRegressor(tuned)","MAE / R2","0.225 / 0.9638","0.205 / 0.9908","+rolling/delta/orbital features, RandomizedSearchCV"),
    ("battery_health","XGBRegressor(tuned)","MAE / R2","0.064 / 0.9992","0.043 / 0.9996",""),
    ("thermal_health","XGBRegressor(tuned)","MAE / R2","0.276 / 0.9865","0.149 / 0.9973",""),
    ("power_health","XGBRegressor(tuned)","MAE / R2","0.042 / 0.9999","0.032 / 0.9999",""),
    ("payload_health","XGBRegressor(tuned)","MAE / R2","0.000 / 1.0000","0.000 / 1.0000","deterministic input, ceiling already hit"),
    ("anomaly_flag","XGBClassifier(weighted+tuned)","Acc / F1","0.9954 / 0.6667","0.9988 / 0.9231","scale_pos_weight=137"),
    ("temperature_anomaly","XGBClassifier(weighted+tuned)","F1","0.800","1.000","scale_pos_weight=265"),
    ("battery_voltage_anomaly","XGBClassifier(weighted+tuned)","F1","0.000","1.000","scale_pos_weight=863, only 4-6 positives total, watch for overfit on so few examples"),
    ("solar_power_anomaly","XGBClassifier(weighted+tuned)","F1","1.000","1.000","scale_pos_weight=431"),
]
for r, row_vals in enumerate(rows, start=1):
    for c, val in enumerate(row_vals, start=1):
        cell = ws2.cell(row=r, column=c, value=val)
        cell.font = header_font if r == 1 else body_font
        if r == 1: cell.fill = header_fill
for c in range(1,7):
    ws2.column_dimensions[get_column_letter(c)].width = 26

wb.save(OUT)
print("Saved:", OUT, "rows:", len(tel_out))